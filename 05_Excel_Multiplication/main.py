# ==================== EXCEL MULTIPLICATION TABLE GENERATOR ====================
# Creates N×N multiplication table and saves as Excel file to Desktop

import openpyxl
import sys
from pathlib import Path
from openpyxl.styles import Font

# ==================== MAIN PROGRAM ====================
if len(sys.argv) == 2:
    number = int(sys.argv[1])

    # Create new workbook and setup sheet
    excelFile = openpyxl.Workbook()
    firstSheet = excelFile.active
    firstSheet.title = "firstSheet"

    # Define bold font for headers
    bold_font = Font(bold=True)
    
    # Add column headers
    for i in range(1, number + 1):
        cell_top = firstSheet.cell(row=1, column=i+1)
        cell_top.value = i
        cell_top.font = bold_font
        
    # Add row headers
    for i in range(1, number + 1):
        cell_side = firstSheet.cell(row=i+1, column=1)
        cell_side.value = i
        cell_side.font = bold_font

    # Fill multiplication table
    for x in range(1, number+1):
        for y in range(1, number+1):
            firstSheet.cell(row=x+1, column=y+1).value = x*y

    # Save to Desktop
    save_path = Path.home() / Path('Desktop') / f'multiplication_table_{str(number)}.xlsx'
    excelFile.save(save_path)

else:
    print("Please enter exactly two arguments: file_name and a number.")