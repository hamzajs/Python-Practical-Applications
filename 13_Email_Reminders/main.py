"""
Send email reminders to members with unpaid monthly dues.

The script reads member payment records from an Excel workbook, identifies
unpaid members, and sends each of them an email reminder for the latest month.
"""

import openpyxl, sys
import smtplib
from pathlib import Path

# Load the member records from the workbook on the Desktop.
try:
    excelFile = openpyxl.load_workbook(Path.home() / Path('Desktop','members.xlsx'))
    Sheet1 = excelFile.active
except FileNotFoundError as E:
    sys.exit("The file of Track monthly payments is Not Found!")

# Collect the email addresses of members whose latest payment is not marked paid.
# Start from the second row to skip the worksheet's column headings.
rows = Sheet1.iter_rows(min_row=2, values_only=True)
unpaidMembers = {}

for row in rows:
    # Read the member's name and email; the final column contains payment status.
    name = row[0]
    email = row[1]
    payment = row[-1]

    # Keep only members who do not have a "paid" status for the latest month.
    if payment != 'paid':
        unpaidMembers[name] = email

# Do not request credentials or connect to the mail server when there is nothing to send.
if not unpaidMembers:
    sys.exit("All of members are Paid.")

# Request the sender's email account credentials.
sender_email = input("Enter Sender Email: ").strip()
password = input("Enter Sender Password: ").strip()

# Use the latest month listed in the worksheet as the reminder period.
latestMonth = Sheet1.cell(row=1, column=Sheet1.max_column).value

# Connect securely to Gmail's SMTP server and authenticate the sender.
server = smtplib.SMTP('smtp.gmail.com', 587)
server.starttls()
server.login(sender_email, password)
print("Login successs")

# Send a personalized payment reminder to each unpaid member.
for name, email in unpaidMembers.items():
    # Personalize the message with the member's name and the latest month.
    text = f"""
Dear {name},

Records show that you have not paid dues for {latestMonth}.
Please make this payment as soon as possible.

Thank you!
    """
    # Add the email subject before the message body.
    message = f"Subject: {latestMonth} dues unpaid.\n{text}"

    # Send the message and report the recipient address.
    server.sendmail(sender_email, email, message.encode('utf-8'))
    print("Email has been sent to ", email)

# Close the SMTP connection after all reminders have been sent.
server.quit()